# -*- coding: utf-8 -*-
"""
AGI-PM 培训资料生成器
输出：
  1) AGI-PM-Training-Manual-VI-ZH-EN.pdf   培训手册（三语，章节式，HTML→Edge 打印 PDF）
  2) AGI-PM-Training-Slides-VI-ZH-EN.pptx  培训 PPT（26 页，每页三语并列：越/中/英）
"""
import os
import subprocess
import sys

sys.stdout.reconfigure(encoding="utf-8")
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

OUT_DIR = r"C:\Users\tomal\OneDrive\Desktop"
MANUAL_HTML = os.path.join(OUT_DIR, "_agi_pm_manual.html")
MANUAL_PDF = os.path.join(OUT_DIR, "AGI-PM-Training-Manual-VI-ZH-EN.pdf")
SLIDES_PPTX = os.path.join(OUT_DIR, "AGI-PM-Training-Slides-VI-ZH-EN.pptx")
EDGE = r"C:\Program Files (x86)\Microsoft\Edge\Application\msedge.exe"

ORANGE = RGBColor(0xE8, 0x72, 0x0C)
DARK = RGBColor(0x2C, 0x3E, 0x50)
GREEN = RGBColor(0x1E, 0x8E, 0x3E)
BLUE = RGBColor(0x2B, 0x6C, 0xB0)
GRAY = RGBColor(0x5F, 0x6B, 0x76)

# ============================================================================
# 内容：每页 = 标题(三语) + 三语要点（顺序：越 → 中 → 英）
# ============================================================================
S = [
 dict(
  t=("AGI-PM — ĐÀO TẠO NGƯỜI DÙNG", "AGI-PM 用户培训手册", "AGI-PM User Training"),
  vi=["Hệ thống quản lý dự án & hoa hồng AGI-GS", "Tài liệu đào tạo chính thức — Ba ngôn ngữ"],
  zh=["AGI-GS 项目与佣金管理系统", "正式培训资料 — 越南语 / 中文 / English 三语对照"],
  en=["AGI-GS Project & Commission Management System", "Official training material — Vietnamese / Chinese / English"],
 ),
 dict(
  t=("NỘI DUNG ĐÀO TẠO", "培训议程", "Training Agenda"),
  vi=["Đăng nhập & chế độ đào tạo · Điều hướng", "CRM — Dự án tiềm năng · WON — Dự án thắng", "LOST — Dự án thất bại · Quy trình phê duyệt", "Người dùng & phân quyền · Ứng dụng di động"],
  zh=["登录与培训模式 · 界面导航", "CRM 潜在项目 · WON 签约项目", "LOST 失败项目 · 审批流程", "用户与授权 · 手机 APP 使用"],
  en=["Login & training mode · Navigation", "CRM — Potential projects · WON — Won projects", "LOST — Failed projects · Approval workflow", "Users & permissions · Mobile APP"],
 ),
 dict(
  t=("TỔNG QUAN HỆ THỐNG", "系统概览", "System Overview"),
  vi=["6 phân module: Dashboard, CRM, WON, LOST, Phê duyệt, Người dùng", "Chạy trên nền web: dùng được trên máy tính & điện thoại", "Dữ liệu lưu trên cloud — máy nào đăng nhập cũng thấy giống nhau", "3 ngôn ngữ: Tiếng Việt / 中文 / English"],
  zh=["6 大模块：看板、CRM、签约、失败、审批、用户与授权", "网页系统：电脑和手机都能用（APP 即网页壳）", "数据存云端：任何设备登录看到的内容一致", "三语支持：越南语 / 中文 / English，随时切换"],
  en=["6 modules: Dashboard, CRM, WON, LOST, Approvals, Users", "Web-based: works on PC & phone (APP = web shell)", "Cloud data: same view on every device", "3 languages: Vietnamese / Chinese / English"],
 ),
 dict(
  t=("ĐĂNG NHẬP & CHẾ ĐỘ ĐÀO TẠO", "登录与培训模式", "Login & Training Mode"),
  vi=["Dùng tên đăng nhập + mật khẩu do quản trị cung cấp", "SAU MỖI LẦN ĐĂNG NHẬP: dữ liệu nghiệp vụ tự động trống", "Nhấn DEMO để nạp dữ liệu mẫu phục vụ luyện tập", "Danh sách người dùng KHÔNG bao giờ bị ảnh hưởng"],
  zh=["使用管理员提供的账号密码登录", "每次登录成功：业务数据自动清空（空场开始）", "点 DEMO 按钮灌入演示数据，用于练习", "用户名录永远不受影响（与 DEMO 彻底脱钩）"],
  en=["Log in with the account provided by the admin", "EVERY login: business data starts EMPTY", "Press DEMO to load sample data for practice", "The user list is NEVER affected by DEMO"],
 ),
 dict(
  t=("ĐIỀU HƯỚNG GIAO DIỆN", "界面导航", "Navigation"),
  vi=["Thanh trên có 6 thẻ: Tổng quan · Tiềm năng · Thắng · Thất bại · Phê duyệt · Người dùng", "Chỉ người có quyền mới thấy thẻ «Người dùng»", "Bên phải: đổi ngôn ngữ, avatar, tên, đăng xuất"],
  zh=["顶栏 6 个标签：看板 · 潜在 · 签约 · 失败 · 审批 · 用户", "「用户」标签仅授权人员可见", "右侧：语言切换、头像、姓名、退出登录"],
  en=["Top bar has 6 tabs: Board · Potential · Won · Failed · Approvals · Users", "The Users tab is only visible to authorized staff", "Right side: language switch, avatar, name, logout"],
 ),
 dict(
  t=("ĐỔI NGÔN NGỮ", "切换语言", "Language Switch"),
  vi=["Nút Lang ở góc phải trên: VI / 中 / EN", "Toàn bộ giao diện đổi ngay, không cần tải lại", "Lựa chọn được ghi nhớ cho lần đăng nhập sau"],
  zh=["右上角 Lang 按钮：VI / 中 / EN", "整站即时切换，无需刷新", "选择会被记住，下次登录保持"],
  en=["Lang button top-right: VI / 中 / EN", "Instant switch, no reload needed", "Your choice is remembered"],
 ),
 dict(
  t=("NÚT DEMO", "DEMO 按钮", "DEMO Button"),
  vi=["Có dữ liệu → nhấn là XÓA SẠCH dữ liệu nghiệp vụ", "Không có dữ liệu → nhấn là NẠP bộ dữ liệu mẫu", "Chỉ tác động 4 bảng nghiệp vụ — người dùng an toàn 100%", "Trong lúc đào tạo: cứ thoải mái nạp/xoá để luyện"],
  zh=["有数据时点击 = 清空全部业务数据", "无数据时点击 = 还原演示数据", "只动 4 张业务表，用户列表 100% 安全", "培训期间可随时反复清空/还原练习"],
  en=["With data → press = CLEAR all business data", "Empty → press = LOAD the demo dataset", "Only touches 4 business tables — users are 100% safe", "Load/clear freely during training"],
 ),
 dict(
  t=("TỔNG QUAN (DASHBOARD)", "看板", "Dashboard"),
  vi=["Thống kê tổng quan: số dự án, giá trị, tỷ lệ thắng/thất", "Thống kê theo từng nhân viên kinh doanh", "Lọc theo thời gian để xem xu hướng"],
  zh=["整体统计：项目数、金额、赢/输率", "按销售人员分组的业绩统计", "支持按时间筛选查看趋势"],
  en=["Overview stats: project counts, value, win/lost rates", "Per-salesperson performance breakdown", "Time filters for trends"],
 ),
 dict(
  t=("CRM — TẠO DỰ ÁN TIỀM NĂNG", "CRM — 新增潜在项目", "CRM — New Potential Project"),
  vi=["Nhấn «+ New CRM» để mở form nhập", "Nhân viên kinh doanh: tên của BẠN được tự điền & khóa (không chọn người khác)", "Quản lý: chọn tự do từ danh sách sales/sales director + tom, cuong, james", "Tên đăng nhập = salesperson: phải là chữ thường & số"],
  zh=["点「+ New CRM」打开录入表单", "销售/销售总监本人录入：销售字段自动填本人并锁死", "经理角色：可从下拉选任意销售/销售总监 + 固定 tom/cuong/james", "销售人员下拉数据源 = 用户名录中的销售岗，自动联动"],
  en=["Press + New CRM to open the entry form", "Sales users: your own name is auto-filled and locked", "Managers: pick anyone from the dropdown (sales roles + tom/cuong/james)", "Dropdown mirrors the user list — always in sync"],
 ),
 dict(
  t=("CRM — TẢI BÁO GIÁ PDF", "CRM — 报价单 PDF 上传", "CRM — Quote PDF Upload"),
  vi=["3 khe: Thiết bị / Lắp đặt / Thiết bị+Lắp đặt", "Chỉ nhận file PDF; tải thành công sẽ lưu vào cloud", "Tải file mới đè lên khe cũ → file cũ tự động xoá (tiết kiệm dung lượng)", "Xoá dự án → các file đính kèm cũng bị xoá theo"],
  zh=["三个槽位：设备 / 安装 / 设备&安装", "仅接受 PDF，上传成功存入云端", "重新上传自动删除被替换的旧文件（省空间）", "删除记录时，附件文件一并捆绑删除"],
  en=["3 slots: Equipment / Installation / Equipment+Installation", "PDF only; stored in cloud after upload", "Re-uploading a slot auto-deletes the replaced file", "Deleting the project also deletes its files"],
 ),
 dict(
  t=("CRM — SỬA / SAO CHÉP / XOÁ", "CRM — 编辑 / 复制 / 删除", "CRM — Edit / Copy / Delete"),
  vi=["«Sửa quyền» chỉnh sửa thông tin dự án", "«Xoá» xoá bản ghi (cần quyền; file đi kèm cũng bị xoá)", "Dữ liệu xoá KHÔNG thể phục hồi — hãy cẩn thận", "Chuyển sang WON/LOST sẽ rời khỏi danh sách tiềm năng"],
  zh=["「编辑」修改项目信息", "「删除」删除记录（需权限，附件一并删除）", "删除不可恢复——请谨慎操作", "转签约/转失败后会离开潜在列表"],
  en=["Edit updates project details", "Delete removes the record (permission required, files too)", "Deletion is permanent — be careful", "Converting to WON/LOST moves it out of the list"],
 ),
 dict(
  t=("WON — TẠO DỰ ÁN THẮNG", "WON — 新增签约项目", "WON — New Won Project"),
  vi=["Hai bên ký kết (AGI-GS / AGI-Customer / GS-Customer)", "Khi có AGI-GS: salesperson cố định «AGI-GS Internal»", "Các trường hợp khác: chọn tự do từ danh sách sales", "Điền giá trị hợp đồng & ngày ký"],
  zh=["签约双方：AGI-GS / AGI-Customer / GS-Customer", "凡涉及 AGI-GS：销售人员锁定为「AGI-GS Internal」", "其他组合：销售下拉可选任意销售", "填写合同金额与签约日期"],
  en=["Signing parties: AGI-GS / AGI-Customer / GS-Customer", "When AGI-GS is a party: salesperson locks to AGI-GS Internal", "Other combos: free sales dropdown", "Enter contract value & signing date"],
 ),
 dict(
  t=("WON — ĐÍNH KÈM & PHỤ LỤC", "WON — 附件与补充协议", "WON — Attachments & Addenda"),
  vi=["Đính kèm hợp đồng lưu trên cloud (R2 file cabinet)", "Phụ lục: tải PDF bổ sung khi thay đổi điều khoản", "Chỉ người có quyền mới xem mục hoa hồng"],
  zh=["合同附件存云端文件柜（R2）", "补充协议：条款变更时上传补充 PDF", "佣金区块仅对有权限的人显示"],
  en=["Contract attachments stored in cloud (R2)", "Addenda: upload extra PDFs for term changes", "Commission block only visible to authorized roles"],
 ),
 dict(
  t=("LOST — DỰ ÁN THẤT BẠI", "LOST — 失败项目", "LOST — Failed Projects"),
  vi=["Ghi nguyên nhân thất bại để rút kinh nghiệm", "Có thể «phục hồi» về danh sách tiềm năng", "Thống kê nguyên nhân giúp cải thiện tỷ lệ thắng"],
  zh=["记录失败原因，沉淀经验", "可一键恢复回潜在项目列表", "失败原因统计帮助提升赢单率"],
  en=["Record failure reasons for learning", "One click to restore back to Potential", "Reason statistics help improve win rate"],
 ),
 dict(
  t=("PHÊ DUYỆT — KHỞI TẠO", "审批流 — 发起", "Approvals — Initiating"),
  vi=["Các hành động nhạy cảm (thêm/sửa/xoá người dùng, tạo dự án…)", "Người không đủ quyền trực tiếp → hệ thống tạo «đơn phê duyệt»", "Chọn người phê duyệt khi gửi đơn", "Cô lập đầy đủ: không ai được bỏ qua quy trình"],
  zh=["敏感操作（增删改用户、建项目等）", "权限不足时自动转为「审批单」，不直接执行", "提交时选择审批人", "闭环设计：任何人都绕不过审批"],
  en=["Sensitive actions (user CRUD, project creation...)", "Without direct permission, an approval request is created", "Pick approvers when submitting", "Closed loop: nobody can bypass the flow"],
 ),
 dict(
  t=("PHÊ DUYỆT — XỬ LÝ", "审批流 — 处理", "Approvals — Processing"),
  vi=["Tom / Cuong / James: bất kỳ ai duyệt là có hiệu lực", "Biểu tượng báo hiệu đơn chờ ở thẻ Phê duyệt", "Xem chi tiết đơn → Duyệt hoặc Từ chối", "Từ chối = hành động bị bỏ qua, dữ liệu không đổi"],
  zh=["tom / cuong / james 任一人审批即生效", "有新审批时导航栏出现红色角标提醒", "查看审批详情 → 通过或拒绝", "拒绝 = 操作作废，数据保持不变"],
  en=["tom / cuong / james — any one approval decides", "A red badge on the Approvals tab shows pending items", "Open details → Approve or Reject", "Reject = action discarded, data unchanged"],
 ),
 dict(
  t=("DANH SÁCH NGƯỜI DÙNG", "用户列表", "User List"),
  vi=["Cột STT để đếm nhanh số lượng người dùng", "3 nút thao tác: Sửa quyền · Sao chép · Xoá", "Trạng thái: Đang làm việc / Nghỉ việc", "Danh sách tách rời hoàn toàn khỏi DEMO"],
  zh=["序号列，一眼数清总人数", "三个操作按钮：编辑权限 · 复制 · 删除", "状态：在职 / 离职", "用户列表与 DEMO 彻底脱钩，永不误删"],
  en=["No. column to count users at a glance", "3 actions: Edit Permission · Copy · Delete", "Status: Active / Departed", "User list is fully decoupled from DEMO"],
 ),
 dict(
  t=("TẠO NGƯỜI DÙNG MỚI", "新增用户", "New User"),
  vi=["Tên đăng nhập: CHỈ chữ thường (a-z) & số (0-9)", "Chữ hoa/tiếng Việt/ký tự → bị chặn ngay khi gõ", "Phân công vị trí: Sales, Giám đốc kinh doanh, Quan sát…", "Người mới cần qua phê duyệt trước khi kích hoạt"],
  zh=["登录账号：只允许小写字母 + 数字", "输入大写/中文/符号会被当场拦截纠正", "选择职位：销售、销售总监、观察者等", "新用户需走审批，通过后账号才激活"],
  en=["Login: lowercase letters & digits ONLY", "Uppercase/local chars are blocked & auto-corrected", "Assign position: Sales, Sales Director, Observer…", "New users go through approval before activation"],
 ),
 dict(
  t=("SAO CHÉP NGƯỜI DÙNG", "复制用户", "Copy User"),
  vi=["Dùng khi nhiều người cùng một vai trò/v quyền", "Nhấn «Sao chép» → form mới đã điền sẵn toàn bộ quyền", "Chỉ cần đổi tên & tên đăng nhập → lưu", "Vẫn đi qua phê duyệt như thêm mới bình thường"],
  zh=["同角色多人时效率神器", "点「复制」→ 新窗口权限全部预填好", "只需改姓名和登录账号 → 保存", "与新增一样走审批闸门，不绕流程"],
  en=["Perfect when many people share one role", "Press Copy → all permissions pre-filled", "Just change name & login → save", "Same approval gate as normal creation"],
 ),
 dict(
  t=("CÁC LỰA CHỌN PHÂN QUYỀN", "权限决策选项", "Permission Decisions"),
  vi=["Mặc định (theo chức vụ) · Ẩn mục này", "Cho phép xem toàn công ty (đặc cách)", "Khởi tạo phê duyệt (hành động phải xin duyệt)", "Sửa / Thêm / Xoá: cho phép hoặc ẩn nút"],
  zh=["保留默认（按职位）· 隐藏该项", "特别授权：显示全公司数据", "发起审批：操作需走审批单", "编辑/新增/删除：允许或隐藏按钮"],
  en=["Keep default (by position) · Hide this item", "Special grant: show whole-company data", "Initiate approval: action goes through a request", "Edit / Add / Delete: allow or hide the button"],
 ),
 dict(
  t=("PHẠM VI DỮ LIỆU", "数据范围", "Data Scope"),
  vi=["«Anh ấy có thể xem tôi» — ai được xem dự án của tôi", "«Tôi có thể xem ai» — tôi được xem dự án của ai", "Quan hệ hai chiều, lưu tự động sau khi lưu form"],
  zh=["「他可见我」—— 谁能看我的项目", "「我可见他」—— 我能看谁的项目", "双向关系，保存表单后自动生效"],
  en=["He-can-see-me: who may view my projects", "I-can-see-them: whose projects I may view", "Two-way relations, saved with the form"],
 ),
 dict(
  t=("VAI TRÒ ĐẶC BIỆT", "特殊角色规则", "Special Roles"),
  vi=["Sales & Giám đốc kinh doanh: 5 mục hoa hồng WON bị khóa «Ẩn»", "Quan sát (Observer): xem mặc định, không sửa/xoá, ẩn hoa hồng", "Quản trị có thể chỉnh nhưng hệ thống luôn bảo vệ khoản hoa hồng"],
  zh=["销售/销售总监：WON 佣金 5 项锁定为「隐藏该项」", "观察者：默认查看权限，无增删改，佣金屏蔽", "管理员可调其他项，但佣金保护始终兜底"],
  en=["Sales & Directors: 5 WON commission items locked to Hide", "Observer: default view rights, no add/edit/delete, no commission", "Admins can adjust others; commission stays protected"],
 ),
 dict(
  t=("QUẢN LÝ MẬT KHẨU", "密码管理", "Password Management"),
  vi=["Đổi mật khẩu ngay trong tab Mật khẩu", "Quên mật khẩu → gửi yêu cầu, người duyệt đặt lại", "Mật khẩu sau khi đặt lại: 66668888 (đổi lại ngay sau khi vào)"],
  zh=["在「密码管理」页签内自行改密", "忘记密码 → 提交申请，审批人重置", "重置后的临时密码：66668888（登录后请立即修改）"],
  en=["Change password in the Password tab", "Forgot? Submit a request; approvers reset it", "Reset password is 66668888 (change it right away)"],
 ),
 dict(
  t=("SỬ DỤNG ỨNG DỤNG DI ĐỘNG", "手机 APP 使用", "Using the Mobile APP"),
  vi=["APP khóa ngang màn hình & chiếm toàn màn hình", "Quét từ mép trên để kéo thanh thông báo xuống", "Nút «Vừa» thu cả bảng vừa màn hình; nhấn đúp vào bảng để trả lại", "Cập nhật trang web: mở lại APP là tự động có bản mới"],
  zh=["APP 固定横屏 + 沉浸式全屏", "从屏幕上边缘下滑可临时唤出信息条", "「适应」按钮缩放整表到一屏；双击表格还原", "网页更新自动生效：重开 APP 即最新版"],
  en=["APP locks landscape & runs fullscreen", "Swipe from the top edge for the status bar", "Fit button scales tables; double-tap to restore", "Web updates arrive automatically on next app open"],
 ),
 dict(
  t=("TÓM TẮT CHẾ ĐỘ ĐÀO TẠO", "培训规则总结", "Training Rules Summary"),
  vi=["Đăng nhập = nghiệp vụ trống → nhấn DEMO để luyện", "DEMO không bao giờ đụng vào danh sách người dùng", "Xoá người dùng chỉ qua phê duyệt", "Đăng xuất KHÔNG xoá dữ liệu — chỉ đăng nhập mới tự động dọn"],
  zh=["登录 = 业务空场 → 点 DEMO 开始练习", "DEMO 永远不碰用户名录", "删用户只能走审批流程", "退出登录不清数据；只有重新登录才自动清场"],
  en=["Login = empty business data → press DEMO to practice", "DEMO never touches the user list", "User deletion only via approval", "Logout keeps data; only a fresh login auto-clears"],
 ),
 dict(
  t=("HỎI ĐÁP & KẾT THÚC", "常见问题与结束", "Q&A & Wrap-up"),
  vi=["Quên mật khẩu? → Yêu cầu đặt lại qua tab Mật khẩu", "Không thấy dữ liệu? → Nhấn DEMO", "Gặp lỗi? → Chụp màn hình & liên hệ quản trị", "Chúc đào tạo thành công!"],
  zh=["忘记密码？→ 密码页签提交重置申请", "看不到数据？→ 点 DEMO 灌入演示数据", "遇到异常？→ 截图联系管理员", "祝培训顺利！"],
  en=["Forgot password? → Reset via Password tab", "No data? → Press DEMO", "Something wrong? → Screenshot & contact admin", "Happy training!"],
 ),
]

# ============================================================================
# 1) PPTX — 每页三语并列（VI 绿 / ZH 橙 / EN 蓝）
# ============================================================================
def build_pptx():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    def add_text(slide, x, y, w, h, text, size, color, bold=False, align=PP_ALIGN.LEFT):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        lines = text.split("\n")
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line
            p.alignment = align
            for r in p.runs:
                r.font.size = Pt(size)
                r.font.color.rgb = color
                r.font.bold = bold
                r.font.name = "Arial"
        return tb

    for idx, item in enumerate(S):
        slide = prs.slides.add_slide(blank)
        # 顶部色带 + 标题
        bar = slide.shapes.add_shape(1, Inches(0), Inches(0), prs.slide_width, Inches(1.0))
        bar.fill.solid(); bar.fill.fore_color.rgb = ORANGE; bar.line.fill.background()
        add_text(slide, 0.45, 0.18, 12.4, 0.7, "  ".join(item["t"]), 26, RGBColor(255, 255, 255), bold=True)
        # 页码
        add_text(slide, 12.5, 7.05, 0.7, 0.35, str(idx + 1), 12, GRAY, align=PP_ALIGN.RIGHT)

        cols = [("VI", GREEN, item["vi"]), ("中文", ORANGE, item["zh"]), ("EN", BLUE, item["en"])]
        x = 0.5
        for label, color, lines in cols:
            # 语言标签
            tag = slide.shapes.add_shape(1, Inches(x), Inches(1.35), Inches(1.0), Inches(0.38))
            tag.fill.solid(); tag.fill.fore_color.rgb = color; tag.line.fill.background()
            tf = tag.text_frame; tf.text = label
            for p in tf.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for r in p.runs:
                    r.font.size = Pt(14); r.font.bold = True; r.font.color.rgb = RGBColor(255, 255, 255)
                    r.font.name = "Arial"
            # 内容
            body = "\n".join("• " + ln for ln in lines)
            add_text(slide, x, 1.95, 3.95, 4.9, body, 15, DARK)
            x += 4.18

    prs.save(SLIDES_PPTX)
    print("PPTX saved:", SLIDES_PPTX)

# ============================================================================
# 2) 手册 HTML → PDF
# ============================================================================
CSS = """
@page { size: A4; margin: 18mm 15mm; }
* { box-sizing: border-box; }
body { font-family: "Segoe UI","Microsoft YaHei","Arial",sans-serif; color:#22313f; margin:0; font-size:11.5pt; line-height:1.55; }
.cover { height:250mm; display:flex; flex-direction:column; justify-content:center; text-align:center;
  background:linear-gradient(135deg,#e8720c,#f08c2f); color:#fff; border-radius:14px; padding:30px; }
.cover h1 { font-size:34pt; margin:0 0 10px; }
.cover .sub { font-size:15pt; opacity:.95; }
.cover .langs { margin-top:26px; font-size:13pt; letter-spacing:2px; }
.toc { margin:26px 0 10px; }
.toc h2 { color:#e8720c; font-size:16pt; border-bottom:2px solid #e8720c; padding-bottom:4px;}
.toc ol { margin:6px 0 0 18px; padding:0; }
.toc li { margin:3px 0; }
.sec { margin:22px 0 8px; page-break-inside:avoid; }
.sec .no { display:inline-block; background:#e8720c; color:#fff; border-radius:6px; padding:1px 10px; font-weight:700; margin-right:8px;}
.sec h2 { display:inline; font-size:14.5pt; color:#2c3e50; }
.lang { margin:7px 0 0; padding:7px 12px; border-left:4px solid #ccc; background:#f7f9fa; border-radius:0 8px 8px 0; }
.lang b.tag { font-size:10pt; border-radius:4px; padding:0 6px; color:#fff; margin-right:6px; }
.vi { border-color:#1e8e3e; } .vi b.tag { background:#1e8e3e; }
.zh { border-color:#e8720c; } .zh b.tag { background:#e8720c; }
.en { border-color:#2b6cb0; } .en b.tag { background:#2b6cb0; }
.lang ul { margin:4px 0 2px 16px; padding:0; }
.lang li { margin:2px 0; }
"""

def build_manual_html():
    parts = []
    parts.append("<html><head><meta charset='utf-8'><style>%s</style></head><body>" % CSS)
    c = S[0]
    parts.append("<div class='cover'><h1>AGI-PM</h1><div class='sub'>%s<br>%s<br>%s</div>" % (
        c["t"][0], c["t"][1], c["t"][2]))
    parts.append("<div class='langs'>Tiếng Việt &nbsp;·&nbsp; 中文 &nbsp;·&nbsp; English</div>")
    parts.append("<div class='sub' style='margin-top:18px'>%s</div></div>" % c["vi"][0])
    parts.append("<div class='toc'><h2>MỤC LỤC · 目录 · Contents</h2><ol>")
    for i, it in enumerate(S[1:], start=1):
        parts.append("<li>%s / %s / %s</li>" % it["t"])
    parts.append("</ol></div>")
    for i, it in enumerate(S[1:], start=1):
        parts.append("<div class='sec'><span class='no'>%d</span><h2>%s · %s · %s</h2></div>" % (i, it["t"][0], it["t"][1], it["t"][2]))
        for cls, tag, key in (("vi", "TIẾNG VIỆT", "vi"), ("zh", "中文", "zh"), ("en", "ENGLISH", "en")):
            parts.append("<div class='lang %s'><b class='tag'>%s</b><ul>" % (cls, tag))
            for ln in it[key]:
                parts.append("<li>%s</li>" % ln)
            parts.append("</ul></div>")
    parts.append("</body></html>")
    with open(MANUAL_HTML, "w", encoding="utf-8") as f:
        f.write("".join(parts))
    print("HTML saved:", MANUAL_HTML)

def html_to_pdf():
    cmd = [EDGE, "--headless", "--disable-gpu", "--no-pdf-header-footer",
           "--print-to-pdf=" + MANUAL_PDF, "file:///" + MANUAL_HTML.replace("\\", "/")]
    r = subprocess.run(cmd, capture_output=True, text=True, timeout=90)
    ok = os.path.exists(MANUAL_PDF) and os.path.getsize(MANUAL_PDF) > 10000
    print("PDF %s (%.1f KB)" % ("saved: " + MANUAL_PDF if ok else "FAILED", 
          os.path.getsize(MANUAL_PDF) / 1024 if os.path.exists(MANUAL_PDF) else 0))
    if not ok:
        print(r.stderr[-600:])

if __name__ == "__main__":
    build_pptx()
    build_manual_html()
    html_to_pdf()
    os.remove(MANUAL_HTML)
    print("ALL DONE")
